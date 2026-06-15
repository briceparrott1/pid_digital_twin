from docx import Document
from docx.oxml.ns import qn
from docx.table import Table


def parse_sop(sop_path: str) -> str:
    doc = Document(sop_path)

    lines = []
    for para in doc.paragraphs:
        if para.text.strip():
            lines.append(para.text.strip())

    # doc.tables only returns tables that are direct children of the body,
    # so tables wrapped in a content control (w:sdt) are skipped. Walk the
    # whole body to catch those too.
    for tbl in doc.element.body.iter(qn("w:tbl")):
        table = Table(tbl, doc)
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    lines.append(cell.text.strip())

    return "\n".join(lines)
